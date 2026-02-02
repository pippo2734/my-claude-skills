#!/usr/bin/env python3
"""
RURA 12月拡大クラブ 導入1ヶ月フォローアップMTG スライド生成スクリプト
ルネサンスブランドに沿ったPowerPointを生成します（視認性改善版）
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ブランドカラー（ルネサンス）
RENAISSANCE_RED = RGBColor(200, 30, 30)  # より見やすい赤
LIGHT_RED = RGBColor(255, 230, 230)  # 薄い赤（背景用）
GRAY_BAND = RGBColor(240, 240, 240)  # 薄いグレー
BLACK = RGBColor(0, 0, 0)
WHITE = RGBColor(255, 255, 255)
DARK_GRAY = RGBColor(80, 80, 80)
BLUE = RGBColor(41, 98, 164)  # アクセント用
LIGHT_BLUE = RGBColor(230, 240, 250)

# スライドサイズ（16:9）
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)


def add_logo_text(slide):
    """右上にルネサンスロゴ（テキスト）を追加"""
    logo_box = slide.shapes.add_textbox(
        Inches(10.3), Inches(0.2), Inches(2.8), Inches(0.5)
    )
    tf = logo_box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = "RENAISSANCE"
    run.font.size = Pt(20)
    run.font.bold = True
    run.font.color.rgb = RENAISSANCE_RED
    run.font.name = "Arial"


def add_header_line(slide):
    """ヘッダー下の赤線を追加"""
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0.75),
        SLIDE_WIDTH, Pt(4)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RENAISSANCE_RED
    line.line.fill.background()


def create_title_slide(prs, title, subtitle=""):
    """タイトルスライド（表紙用）を作成"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    add_logo_text(slide)

    # 中央のグレー帯
    band = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(2.5),
        SLIDE_WIDTH, Inches(2.8)
    )
    band.fill.solid()
    band.fill.fore_color.rgb = GRAY_BAND
    band.line.fill.background()

    # タイトルテキスト
    title_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(2.7), Inches(12.3), Inches(2.0)
    )
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = title
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = BLACK
    p.font.name = "Meiryo"

    if subtitle:
        sub_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(5.5), Inches(12.3), Inches(0.8)
        )
        tf = sub_box.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = subtitle
        p.font.size = Pt(24)
        p.font.color.rgb = DARK_GRAY
        p.font.name = "Meiryo"

    return slide


def create_section_slide(prs, section_number, section_title):
    """セクション区切りスライドを作成（大きな数字付き）"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    add_logo_text(slide)

    # 左側に大きな数字の円
    circle = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(1.5), Inches(2.2),
        Inches(2.5), Inches(2.5)
    )
    circle.fill.solid()
    circle.fill.fore_color.rgb = RENAISSANCE_RED
    circle.line.fill.background()

    # 数字
    num_box = slide.shapes.add_textbox(
        Inches(1.5), Inches(2.5), Inches(2.5), Inches(2.0)
    )
    tf = num_box.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = str(section_number)
    p.font.size = Pt(72)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Arial"

    # セクションタイトル
    title_box = slide.shapes.add_textbox(
        Inches(4.5), Inches(2.8), Inches(8), Inches(1.8)
    )
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = section_title
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = BLACK
    p.font.name = "Meiryo"

    return slide


def create_content_slide_with_boxes(prs, title, items):
    """ボックス付きコンテンツスライド"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    add_logo_text(slide)
    add_header_line(slide)

    # タイトル
    title_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(0.15), Inches(9.5), Inches(0.6)
    )
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = RENAISSANCE_RED
    p.font.name = "Meiryo"

    # コンテンツボックスを配置
    y_pos = 1.1
    for item in items:
        if isinstance(item, dict) and 'heading' in item:
            # 見出し付きボックス
            box = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(0.5), Inches(y_pos),
                Inches(12.3), Inches(1.3)
            )
            box.fill.solid()
            box.fill.fore_color.rgb = LIGHT_RED if item.get('highlight') else GRAY_BAND
            box.line.color.rgb = RENAISSANCE_RED if item.get('highlight') else RGBColor(200, 200, 200)

            # 見出し
            head_box = slide.shapes.add_textbox(
                Inches(0.8), Inches(y_pos + 0.15), Inches(11.5), Inches(0.5)
            )
            tf = head_box.text_frame
            p = tf.paragraphs[0]
            p.text = item['heading']
            p.font.size = Pt(22)
            p.font.bold = True
            p.font.color.rgb = RENAISSANCE_RED
            p.font.name = "Meiryo"

            # 内容
            if item.get('content'):
                cont_box = slide.shapes.add_textbox(
                    Inches(0.8), Inches(y_pos + 0.6), Inches(11.5), Inches(0.6)
                )
                tf = cont_box.text_frame
                p = tf.paragraphs[0]
                p.text = item['content']
                p.font.size = Pt(20)
                p.font.color.rgb = BLACK
                p.font.name = "Meiryo"

            y_pos += 1.5

    return slide


def create_list_slide(prs, title, items, use_numbers=False):
    """リスト形式のスライド（大きなフォント）"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    add_logo_text(slide)
    add_header_line(slide)

    # タイトル
    title_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(0.15), Inches(9.5), Inches(0.6)
    )
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = RENAISSANCE_RED
    p.font.name = "Meiryo"

    # リストアイテム
    y_pos = 1.2
    for i, item in enumerate(items, 1):
        if use_numbers:
            # 番号の円
            circle = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                Inches(0.6), Inches(y_pos),
                Inches(0.6), Inches(0.6)
            )
            circle.fill.solid()
            circle.fill.fore_color.rgb = RENAISSANCE_RED
            circle.line.fill.background()

            num_box = slide.shapes.add_textbox(
                Inches(0.6), Inches(y_pos + 0.08), Inches(0.6), Inches(0.5)
            )
            tf = num_box.text_frame
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.text = str(i)
            p.font.size = Pt(22)
            p.font.bold = True
            p.font.color.rgb = WHITE
            p.font.name = "Arial"

            text_left = 1.4
        else:
            # 四角のブレット
            bullet = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(0.7), Inches(y_pos + 0.2),
                Inches(0.2), Inches(0.2)
            )
            bullet.fill.solid()
            bullet.fill.fore_color.rgb = RENAISSANCE_RED
            bullet.line.fill.background()
            text_left = 1.1

        # テキスト
        item_box = slide.shapes.add_textbox(
            Inches(text_left), Inches(y_pos + 0.05), Inches(11), Inches(0.6)
        )
        tf = item_box.text_frame
        p = tf.paragraphs[0]
        p.text = item
        p.font.size = Pt(24)
        p.font.color.rgb = BLACK
        p.font.name = "Meiryo"

        y_pos += 1.0

    return slide


def create_two_column_slide(prs, title, left_items, right_items, left_title="", right_title=""):
    """2カラムスライド"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    add_logo_text(slide)
    add_header_line(slide)

    # タイトル
    title_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(0.15), Inches(9.5), Inches(0.6)
    )
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = RENAISSANCE_RED
    p.font.name = "Meiryo"

    # 左カラムボックス
    left_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.5), Inches(1.1),
        Inches(6), Inches(5.8)
    )
    left_box.fill.solid()
    left_box.fill.fore_color.rgb = GRAY_BAND
    left_box.line.color.rgb = RGBColor(200, 200, 200)

    # 左カラムタイトル
    if left_title:
        lt_box = slide.shapes.add_textbox(
            Inches(0.8), Inches(1.3), Inches(5.4), Inches(0.5)
        )
        tf = lt_box.text_frame
        p = tf.paragraphs[0]
        p.text = left_title
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = RENAISSANCE_RED
        p.font.name = "Meiryo"

    # 左カラム内容
    y_pos = 2.0 if left_title else 1.4
    for item in left_items:
        item_box = slide.shapes.add_textbox(
            Inches(0.8), Inches(y_pos), Inches(5.4), Inches(0.6)
        )
        tf = item_box.text_frame
        p = tf.paragraphs[0]
        p.text = f"• {item}"
        p.font.size = Pt(20)
        p.font.color.rgb = BLACK
        p.font.name = "Meiryo"
        y_pos += 0.7

    # 右カラムボックス
    right_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(6.8), Inches(1.1),
        Inches(6), Inches(5.8)
    )
    right_box.fill.solid()
    right_box.fill.fore_color.rgb = LIGHT_BLUE
    right_box.line.color.rgb = BLUE

    # 右カラムタイトル
    if right_title:
        rt_box = slide.shapes.add_textbox(
            Inches(7.1), Inches(1.3), Inches(5.4), Inches(0.5)
        )
        tf = rt_box.text_frame
        p = tf.paragraphs[0]
        p.text = right_title
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = BLUE
        p.font.name = "Meiryo"

    # 右カラム内容
    y_pos = 2.0 if right_title else 1.4
    for item in right_items:
        item_box = slide.shapes.add_textbox(
            Inches(7.1), Inches(y_pos), Inches(5.4), Inches(0.6)
        )
        tf = item_box.text_frame
        p = tf.paragraphs[0]
        p.text = f"• {item}"
        p.font.size = Pt(20)
        p.font.color.rgb = BLACK
        p.font.name = "Meiryo"
        y_pos += 0.7

    return slide


def create_highlight_slide(prs, title, main_point, sub_points=None):
    """強調ポイントスライド"""
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    add_logo_text(slide)
    add_header_line(slide)

    # タイトル
    title_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(0.15), Inches(9.5), Inches(0.6)
    )
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = RENAISSANCE_RED
    p.font.name = "Meiryo"

    # メインポイントの強調ボックス
    main_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(1), Inches(1.5),
        Inches(11.3), Inches(2.2)
    )
    main_box.fill.solid()
    main_box.fill.fore_color.rgb = RENAISSANCE_RED
    main_box.line.fill.background()

    main_text = slide.shapes.add_textbox(
        Inches(1.3), Inches(2.0), Inches(10.7), Inches(1.5)
    )
    tf = main_text.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = main_point
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Meiryo"

    # サブポイント
    if sub_points:
        y_pos = 4.0
        for point in sub_points:
            point_box = slide.shapes.add_textbox(
                Inches(1.5), Inches(y_pos), Inches(10.3), Inches(0.7)
            )
            tf = point_box.text_frame
            p = tf.paragraphs[0]
            p.text = f"✓ {point}"
            p.font.size = Pt(22)
            p.font.color.rgb = DARK_GRAY
            p.font.name = "Meiryo"
            y_pos += 0.8

    return slide


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    # ===== スライド0: 表紙 =====
    create_title_slide(
        prs,
        "RURA 12月拡大クラブ\n導入1ヶ月フォローアップMTG",
        "運用状況の確認と標準化に向けたステップ"
    )

    # ===== スライド1: 本MTGの位置付け =====
    create_content_slide_with_boxes(
        prs,
        "1. 本MTGの位置付け",
        [
            {"heading": "コンセプト", "content": "共創スタンスでの情報共有・意見交換の場"},
            {"heading": "目的① STEP0の細分化・確認", "content": "キックオフ時の内容を実運用レベルで再確認"},
            {"heading": "目的②「クラブの声を聴く会」", "content": "現場の課題や要望を吸い上げる"},
            {"heading": "目的③ 標準とフィードバック", "content": "本部の標準オペレーション提示と、現場からのFBを両輪で回す"},
        ]
    )

    # ===== スライド2: 本日のゴール =====
    create_highlight_slide(
        prs,
        "2. 本日のゴール（終了時の状態）",
        "「2月の行動計画（シフト修正案含む）」が提出され、\nコミットメントが得られている状態",
        [
            "シフト適正化の履行：「適正シフト」の定義を理解し、次月以降の作成フロー変更を合意",
            "標準ルールの適用：統一ルール（1/15決定事項）の認識ズレを解消",
        ]
    )

    # ===== スライド3: アジェンダ =====
    create_list_slide(
        prs,
        "3. アジェンダ",
        [
            "RURA導入ステップ（STEP0）の確認",
            "標準オペレーションの徹底",
            "現場運用の確認事項",
            "帳票格納ルールの標準化",
            "クラブの声を聴く会・質疑応答",
        ],
        use_numbers=True
    )

    # ===== スライド4: セクション【1】 =====
    create_section_slide(prs, 1, "RURA導入ステップ\n（STEP0）の確認")

    # ===== スライド5: RURA導入ステップ概要 =====
    create_content_slide_with_boxes(
        prs,
        "【1】RURA導入ステップ（STEP0）の確認",
        [
            {"heading": "テーマ", "content": "運用の土台となる「誘導」と「時間の使い方」についての認識合わせ"},
            {"heading": "ポイント", "content": "キックオフ時のSTEP0をより解像度を上げて確認します", "highlight": True},
        ]
    )

    # ===== スライド6: 誘導率の振り返り =====
    create_two_column_slide(
        prs,
        "1-1. 誘導率の振り返り",
        ["12月誘導率の実績", "1月誘導率の進捗"],
        ["数値に基づいた現状把握", "改善ポイントの特定"],
        "現状共有", "メッセージ"
    )

    # ===== スライド7: 適正シフトの考え方 =====
    create_content_slide_with_boxes(
        prs,
        "1-2. 適正シフトの考え方（STEP0再確認）",
        [
            {"heading": "定義の確認", "content": "「顧客接点時間」と「非顧客接点時間」の明確化"},
            {"heading": "シフト作成ガイドライン", "content": "ガイドラインに沿ったシフト組みができているか？", "highlight": True},
        ]
    )

    # ===== スライド8: 具体的運用基準（シフト） =====
    create_highlight_slide(
        prs,
        "1-3. 具体的運用基準（シフト）",
        "在籍5,500名までのクラブは「受」１ライン体制",
        [
            "現場から「厳しい」等の声がある場合 → セーフィー社による再調査を実施",
            "7月調査結果も参照可能",
            "【重要】2月シフトより上記基準を反映してください",
        ]
    )

    # ===== スライド9: セクション【2】 =====
    create_section_slide(prs, 2, "標準オペレーション")

    # ===== スライド10: 標準オペレーション概要 =====
    create_content_slide_with_boxes(
        prs,
        "【2】標準オペレーション",
        [
            {"heading": "テーマ", "content": "早急に対応が必要な実務タスクと期限の指示"},
            {"heading": "重要度", "content": "必須（Must）項目の確認", "highlight": True},
        ]
    )

    # ===== スライド11: スクール人数メンテナンス =====
    create_highlight_slide(
        prs,
        "2-1. スクール人数メンテナンス",
        "【必達期限】1月末までに「3月クラス」のメンテナンスを完了",
        [
            "対象範囲拡大：テニス・トライネーションも含めて実施を徹底",
            "現状の課題：1月10日時点で未完了のクラブあり",
        ]
    )

    # ===== スライド12: テニスの運用ルール =====
    create_content_slide_with_boxes(
        prs,
        "2-2. テニスの運用ルール",
        [
            {"heading": "ルール", "content": "プランニングシートの持参必須", "highlight": True},
            {"heading": "詳細", "content": "進級の有無に関わらず、必ず持参が必要となります"},
        ]
    )

    # ===== スライド13: セクション【3】 =====
    create_section_slide(prs, 3, "現場運用の確認事項")

    # ===== スライド14: 現場運用概要 =====
    create_content_slide_with_boxes(
        prs,
        "【3】現場運用の確認事項",
        [
            {"heading": "テーマ", "content": "運用開始後におろそかになりがちな情報のキャッチアップとデータ収集の是正"},
        ]
    )

    # ===== スライド15: 情報連携と機器設定 =====
    create_two_column_slide(
        prs,
        "3-1. 情報連携と機器設定",
        ["1勤務1回の確認を徹底", "情報連携ミスの防止"],
        ["必ず「有線接続」で使用", "正しい音量・接続設定の確認"],
        "共有メモの確認", "ネックスピーカー設定"
    )

    # ===== スライド16: データ集約・その他 =====
    create_list_slide(
        prs,
        "3-2. データ集約・その他",
        [
            "退会理由の集約：正しい集約方法ができているか再確認",
            "掲示板登録：登録漏れがないか確認",
        ]
    )

    # ===== スライド17: 帳票格納ルール =====
    create_highlight_slide(
        prs,
        "【4】帳票格納ルールの標準化",
        "【格納スケジュール】毎月15日格納の徹底",
        [
            "目的：事務管理コストを下げるための統一ルール運用",
            "追加格納や修正が発生した場合の手順を確認",
        ]
    )

    # ===== スライド18: クラブの声を聴く会 =====
    create_list_slide(
        prs,
        "【5】クラブの声を聴く会・質疑応答",
        [
            "現場からの質問、要望",
            "発生しているトラブル事例",
            "運用上の課題や改善提案",
        ]
    )

    # ===== スライド19: まとめ =====
    create_highlight_slide(
        prs,
        "まとめ・次回アクション",
        "本日の決定事項を確認し、次回までに実行",
        [
            "シフト修正：2月シフトより反映",
            "メンテナンス期限：1月末まで",
            "帳票格納：毎月15日",
        ]
    )

    # 保存
    output_path = "/home/user/my-claude-skills/RURA_followup_mtg.pptx"
    prs.save(output_path)
    print(f"プレゼンテーションを保存しました: {output_path}")
    return output_path


if __name__ == "__main__":
    main()
